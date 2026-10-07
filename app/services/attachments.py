import io
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader


class AttachmentError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx}


def extract_text(filename: str, data: bytes) -> str:
    extension = PurePath(filename).suffix.lower()
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise AttachmentError(
            f"Formato no soportado en '{filename}': solo PDF y Word (.docx)", status_code=415
        )

    try:
        text = extractor(data).strip()
    except Exception as exc:
        raise AttachmentError(f"No se pudo leer '{filename}': {exc}") from exc

    if not text:
        raise AttachmentError(f"'{filename}' no contiene texto extraible")
    return text


def append_attachments(transcript: str, attachments: list[tuple[str, str]]) -> str:
    """Concatena al transcript el texto de cada adjunto con un separador claro."""
    blocks = [transcript] + [f"--- attachment: {name} ---\n{text}" for name, text in attachments]
    return "\n\n".join(blocks)
